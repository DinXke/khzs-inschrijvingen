"""Excel-overzicht van de goedgekeurde inschrijvingen."""
import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .regels import toon_tijd


def _breedtes(ws):
    for kol in ws.columns:
        ws.column_dimensions[get_column_letter(kol[0].column)].width = max(len(str(c.value or "")) for c in kol) + 2


def excel_overzicht(wedstrijd, inschrijvingen) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inschrijvingen"
    ws.append(["Sessie", "Nr", "Nummer", "Naam", "Voornaam", "Geboortedatum", "Licentie", "Inschrijftijd", "Baan", "Kost (€)"])
    per_zwemmer = defaultdict(lambda: [0, 0.0])
    for i in sorted(inschrijvingen, key=lambda i: (i.nummer.sessie.nummer, i.nummer.volgorde, i.zwemmer.naam)):
        z, n = i.zwemmer, i.nummer
        ws.append([n.sessie.nummer, n.nummer, n.omschrijving, z.naam, z.voornaam, z.geboortedatum,
                   z.licentie, toon_tijd(i.inschrijftijd), i.inschrijfbaan, n.kost])
        per_zwemmer[z][0] += 1
        per_zwemmer[z][1] += n.kost

    ws2 = wb.create_sheet("Per zwemmer")
    ws2.append(["Naam", "Voornaam", "Aantal nummers", "Totaal (€)"])
    for z, (aantal, kost) in sorted(per_zwemmer.items(), key=lambda kv: (kv[0].naam, kv[0].voornaam)):
        ws2.append([z.naam, z.voornaam, aantal, kost])
    ws2.append(["Totaal", "", sum(v[0] for v in per_zwemmer.values()), sum(v[1] for v in per_zwemmer.values())])

    for blad in (ws, ws2):
        for c in blad[1]:
            c.font = Font(bold=True)
        blad.freeze_panes = "A2"
        _breedtes(blad)
    for rij in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        rij[0].number_format = "DD/MM/YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
